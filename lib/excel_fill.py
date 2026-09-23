# -*- coding: utf-8 -*-
"""
excel_fill.py
รับข้อมูลที่กรอกจากฟอร์ม แล้วเขียนลงไฟล์ Excel แม่แบบตาม TERMINAL ที่เลือก
- แต่ละ TERMINAL มีไฟล์ผลลัพธ์ / ชีต / ตำแหน่งเซลล์ ของตัวเอง (ดู TERMINALS ข้างล่าง)
- รัน No. อัตโนมัติเมื่อมีหลายตู้
- คอลัมน์ LINE => "HAL"
- เฉพาะ A3 (C1,C2): คอลัมน์ PAYMENT TAEM => "CASH"

ทั้งหมดทำงานด้วย openpyxl (ไม่มี AI)
"""

from __future__ import annotations

import os
import re
from datetime import datetime
from typing import Dict, List

import openpyxl
from openpyxl.utils import column_index_from_string, get_column_letter

from .xls_convert import ensure_xlsx

LINE_CODE = "HAL"
AGENT_NAME = "HEUNG-A"

# ---------------------------------------------------------------------------
# ตารางแม็ปข้อมูลของแต่ละ TERMINAL
#   header : ตำแหน่งเซลล์ของ Vessel / Voy. (เขียนครั้งเดียว)
#   table  : start_row = แถวแรกของข้อมูลตู้, max_rows = จำนวนแถวที่รองรับ
#            cols = ชื่อฟิลด์ -> ตัวอักษรคอลัมน์
#   ฟิลด์ที่ใช้ได้ใน cols:
#     no, container, booking, size, type, size_combo, pod, status,
#     line, shipper, seal, vessel, voy, vessel_voy, agent, vgm, terminal,
#     commodity, temp, humidity, vent, dg_flag, un_number, remark,
#     over_height, over_width
# ---------------------------------------------------------------------------
# ท่าที่ให้เขียนช่อง REMARK เป็นชื่อบริษัท (Shipper) อย่างเดียว (ข้อมูลผู้ติดต่อ/เวลาส่ง
# ย้ายไปมีช่องของตัวเองในแม่แบบแล้ว ไม่ต้องพ่วงมากับ REMARK อีกต่อไป)
REMARK_SHIPPER_ONLY_KEYS = {"A0", "B3"}

TERMINALS: Dict[str, Dict] = {
    # ---------- A0 : LCMT / LCB1 ----------
    "LCMT Company LTD, ( under LCB1 Group)  A0": {
        "key": "A0",
        "src": "input/A0-SHORE.xls",
        "out": "A0-SHORE.xlsx",
        "sheet": "Sheet1",
        "header": {
            "vessel": "F6", "voy": "J6",
            "submitted_at": "D36",   # ใต้ป้าย "Date / Time :" ที่ A36
            "contact_name": "D37",   # ใต้ป้าย "NAME / ชื่อ :" ที่ A37
            "contact_phone": "D38",  # ใต้ป้าย "TEL / โทร :" ที่ A38
            "contact_email": "D39",  # ใต้ป้าย "EMAIL :" ที่ A39
        },
        "table": {
            "start_row": 14,
            "max_rows": 19,  # แถว 14-32 (แถว 33 เป็นข้อความ disclaimer)
            "cols": {
                "no": "A", "container": "C", "booking": "E",
                "size": "F", "type": "G", "pod": "H",
                "temp": "I", "vent": "J", "remark": "L",
            },
        },
    },
    # ---------- B3 : ESCO ----------
    "ESCO (EASTERN SEA LCH CNTR TML/B3)": {
        "key": "B3",
        "src": "input/B3-SHORE.xls",
        "out": "B3-SHORE.xlsx",
        "sheet": "CHORE CY",
        # C11 เป็นเซลล์ที่ผสาน Vessel+Voy ไว้ด้วยกัน (แม่แบบรุ่นล่าสุด) จึงรวมเป็นช่องเดียว
        "header": {
            "vessel_voy": "C11", "shipper": "C12",
            "submitted_at": "C32",   # ใต้ป้าย "Date / Time" ที่ A32
            "contact_name": "C33",   # ใต้ป้าย "NAME / ชื่อ :" ที่ A33
            "contact_phone": "C34",  # ใต้ป้าย "TEL / โทร :" ที่ A34
            "contact_email": "C35",  # ใต้ป้าย "EMAIL :" ที่ A35
        },
        "table": {
            "start_row": 17,
            "max_rows": 20,
            "cols": {
                "no": "A", "container": "B", "type": "C", "size": "D",
                "pod": "E", "booking": "J",
                "temp": "G", "dg_flag": "H", "remark": "L",
            },
        },
    },
    # ---------- B5/C3 : LCIT ----------
    # เขียนลงชีต "DataImport" (รูปแบบไฟล์นำเข้าจริงของท่าเรือ) ไม่ใช่ "Sheet1" (แบบฟอร์มเดิม)
    "B5/C3 LCIT (LAEM CHABANG INTERNATIONAL TERMINAL CO., LTD)": {
        "key": "B5C3",
        "src": "input/B5C3-SHORE.xls",
        "out": "B5C3-SHORE.xlsx",
        "sheet": "DataImport",
        "header": {
            "submitted_at": "B22",   # ใต้ป้าย "Date / Time :" ที่ A22
            "contact_name": "B23",   # ใต้ป้าย "NAME / ชื่อ :" ที่ A23
            "contact_phone": "B24",  # ใต้ป้าย "TEL / โทร :" ที่ A24
            "contact_email": "B25",  # ใต้ป้าย "EMAIL :" ที่ A25
        },
        "table": {
            "start_row": 2,
            "max_rows": 500,
            "cols": {
                "vessel_voy": "A",   # Vessel Visit
                "opr": "B",          # Opr
                "owner": "C",        # Owner
                "ss": "D",           # SS
                "status": "E",       # Status (F/E)
                "container": "F",    # Cntr No
                "size": "G",         # Size
                "type": "H",         # Type
                "pod": "J",          # POD1
                "booking": "M",      # Bkg No
                "shipper": "N",      # Shipper
                "wt_uom": "P",       # Wt UOM
                "org": "Q",          # ORG
                "remark": "U",       # Remark
            },
            "row_constants": {
                "opr": "HAS", "owner": "HAS", "ss": "EX", "wt_uom": "KG", "org": "LCB",
                "remark": "CASH",
            },
        },
    },
    # ---------- A3 (C1,C2) : Hutchison (HLT) ----------
    "A3 (C1,C2) (Hutchison Laemchabang Terminal Limited, HLT)": {
        "key": "A3C1C2",
        "src": "input/A3C1C2-HUTCHISON.xls",
        "out": "A3C1C2-HUTCHISON.xlsx",
        "sheet": "HPT",
        "header": {
            "submitted_at": "C32",   # ใต้ป้าย "Date / Time" ที่ A32
            "contact_name": "C33",   # ใต้ป้าย "NAME / ชื่อ :" ที่ A33
            "contact_phone": "C34",  # ใต้ป้าย "TEL / โทร :" ที่ A34
            "contact_email": "C35",  # ใต้ป้าย "EMAIL :" ที่ A35
        },
        "table": {
            "start_row": 10,
            "max_rows": 19,
            "cols": {
                "no": "A", "line": "B", "shipper": "C", "size": "D", "type": "E",
                "container": "F", "vessel": "G", "voy": "H", "pod": "I",
                "booking": "K", "seal": "L", "status": "M", "payment": "U",
                "commodity": "N", "temp": "O", "vent": "P",
                "dg_flag": "R", "un_number": "S",
                # REMARK (V) ไม่ต้องเติมข้อมูลแล้ว ปล่อยว่างไว้ตามที่ระบุ
            },
            "row_constants": {"payment": "CASH"},
        },
    },
}

SIZE_SPLIT_RE = re.compile(r"^\s*(\d{2})\s*([A-Za-z/]{1,4})?\s*$")


def split_size(raw: str):
    """'20 GP' -> ('20', 'GP') ; '40HQ' -> ('40', 'HQ')"""
    if not raw:
        return "", ""
    m = SIZE_SPLIT_RE.match(str(raw).replace("  ", " "))
    if m:
        return m.group(1), (m.group(2) or "").upper()
    parts = str(raw).split()
    if len(parts) == 2:
        return parts[0], parts[1].upper()
    return str(raw), ""


def status_code(raw: str) -> str:
    """'FULL' -> 'F' ; 'EMPTY' -> 'E' (สำหรับคอลัมน์ที่ชื่อ STATUS F/E)"""
    s = str(raw or "").strip().upper()
    if s.startswith("F"):
        return "F"
    if s.startswith("E"):
        return "E"
    return s


def _unmerge_anchor(ws, coord: str) -> str:
    """ถ้า coord อยู่ในช่วง merge แต่ไม่ใช่มุมซ้ายบน ให้คืน coord ของมุมซ้ายบนแทน"""
    for rng in ws.merged_cells.ranges:
        if coord in rng:
            return rng.coord.split(":")[0]
    return coord


def _set(ws, coord: str, value) -> None:
    if value is None or value == "":
        return
    ws[_unmerge_anchor(ws, coord)] = value


def _cell(col_letter: str, row: int) -> str:
    return f"{col_letter}{row}"


def build_record_rows(payload: Dict) -> List[Dict]:
    """แปลง payload จากฟอร์มเป็นรายการแถว (หนึ่งแถวต่อหนึ่งตู้)"""
    vessel = (payload.get("vessel") or "").strip()
    voy = (payload.get("voy") or "").strip()
    shipper = (payload.get("shipper") or "").strip()
    pod = (payload.get("pod") or "").strip()
    booking = (payload.get("booking") or "").strip()
    terminal = (payload.get("terminal") or "").strip()

    special_cond = (payload.get("remark") or "").strip()
    contact_name = (payload.get("contactName") or "").strip()
    contact_phone = (payload.get("contactPhone") or "").strip()
    contact_email = (payload.get("contactEmail") or "").strip()
    contact_line = "Contact: " + " ".join(p for p in (contact_name, contact_phone) if p)
    if contact_email:
        contact_line += " " + contact_email

    terminal_key = (TERMINALS.get(terminal) or {}).get("key")
    shipper_only_remark = terminal_key in REMARK_SHIPPER_ONLY_KEYS

    rows: List[Dict] = []
    for item in payload.get("rows", []):
        raw_size = (item.get("size") or "").strip()
        size_num, size_type = split_size(raw_size)
        dg_un = (item.get("dgUn") or "").strip()
        over_height = (item.get("overHeight") or "").strip()
        over_width = (item.get("overWidth") or "").strip()

        if shipper_only_remark:
            remark = shipper
        else:
            remark_parts = [p for p in (special_cond,) if p]
            oversize_bits = []
            if over_height:
                oversize_bits.append(f"Over Height: {over_height}")
            if over_width:
                oversize_bits.append(f"Over Width: {over_width}")
            if oversize_bits:
                remark_parts.append(", ".join(oversize_bits))
            remark_parts.append(contact_line)
            remark = "\n".join(remark_parts)

        rows.append({
            "container": (item.get("container") or "").strip().upper(),
            "size": size_num,
            "type": size_type,
            "size_combo": raw_size.replace(" ", ""),
            "size_raw": raw_size,
            "status": (item.get("status") or "").strip().upper(),
            "vessel": vessel,
            "voy": voy,
            "vessel_voy": f"{vessel} V.{voy}".strip(" V."),
            "shipper": shipper,
            "pod": pod,
            "booking": booking,
            "terminal": terminal,
            "agent": AGENT_NAME,
            "line": LINE_CODE,
            "seal": (item.get("seal") or "").strip(),
            "vgm": (item.get("vgm") or "").strip(),
            "commodity": (item.get("commodity") or "").strip(),
            "temp": (item.get("temp") or "").strip(),
            "humidity": (item.get("humidity") or "").strip(),
            "vent": (item.get("vent") or "").strip(),
            "over_height": over_height,
            "over_width": over_width,
            "un_number": dg_un,
            "dg_flag": "Y" if dg_un else "",
            "remark": remark,
        })
    return rows


def fill(payload: Dict, base_dir: str, cache_dir: str, out_dir: str) -> Dict:
    """
    เขียนข้อมูลลงแม่แบบของ TERMINAL ที่เลือก
    คืน: {"ok":bool, "out_file":str, "rows":int, "error":str, "warnings":[...]}
    """
    terminal = (payload.get("terminal") or "").strip()
    if terminal not in TERMINALS:
        return {"ok": False, "error": f"ไม่รู้จัก TERMINAL: {terminal!r}", "warnings": []}

    conf = TERMINALS[terminal]
    rows = build_record_rows(payload)
    if not rows:
        return {"ok": False, "error": "ยังไม่มีข้อมูลตู้ (rows ว่าง)", "warnings": []}

    warnings: List[str] = []

    src_path = os.path.join(base_dir, conf["src"])
    if not os.path.isfile(src_path):
        return {"ok": False, "error": f"ไม่พบไฟล์แม่แบบ {conf['src']}", "warnings": []}

    template_xlsx = ensure_xlsx(src_path, cache_dir)
    wb = openpyxl.load_workbook(template_xlsx)

    sheet_name = conf["sheet"]
    ws = wb[sheet_name] if sheet_name in wb.sheetnames else wb.active
    if sheet_name not in wb.sheetnames:
        warnings.append(f"ไม่พบชีต {sheet_name!r} ใช้ชีต {ws.title!r} แทน")

    # ---- ส่วนหัว: Vessel / Voy. / Shipper / POD / Booking / ผู้ติดต่อ / เวลาที่ส่ง ----
    # (vessel_voy = ช่องที่ผสาน Vessel+Voy ไว้ด้วยกัน เขียนเป็น "M.V. {vessel} V.{voy}")
    header = conf.get("header", {})
    r0 = rows[0]
    contact_name = (payload.get("contactName") or "").strip()
    contact_phone = (payload.get("contactPhone") or "").strip()
    contact_email = (payload.get("contactEmail") or "").strip()
    submitted_dt = datetime.now().strftime("%d/%m/%Y %H:%M")

    if header.get("vessel"):
        _set(ws, header["vessel"], r0["vessel"])
    if header.get("voy"):
        _set(ws, header["voy"], r0["voy"])
    if header.get("shipper"):
        _set(ws, header["shipper"], r0["shipper"])
    if header.get("pod"):
        _set(ws, header["pod"], r0["pod"])
    if header.get("booking"):
        _set(ws, header["booking"], r0["booking"])
    if header.get("vessel_voy"):
        _set(ws, header["vessel_voy"], f"M.V. {r0['vessel']} V.{r0['voy']}".strip())
    if header.get("contact_name"):
        _set(ws, header["contact_name"], contact_name)
    if header.get("contact_phone"):
        _set(ws, header["contact_phone"], contact_phone)
    if header.get("contact_email"):
        _set(ws, header["contact_email"], contact_email)
    if header.get("contact_combined"):
        _set(ws, header["contact_combined"], " / ".join(p for p in (contact_name, contact_phone) if p))
    if header.get("contact_all"):
        _set(ws, header["contact_all"], " / ".join(p for p in (contact_name, contact_phone, contact_email) if p))
    if header.get("submitted_at"):
        _set(ws, header["submitted_at"], submitted_dt)
    if header.get("submitted_at_by"):
        by = f" - {contact_name}" if contact_name else ""
        _set(ws, header["submitted_at_by"], f"{submitted_dt}{by}")

    tbl = conf["table"]
    start = tbl["start_row"]
    max_rows = tbl["max_rows"]
    cols = tbl["cols"]
    row_constants = tbl.get("row_constants", {})

    # ---- ล้างข้อมูลตัวอย่างเดิม (เฉพาะที่กำหนด clear_first) ----
    # ล้างเฉพาะแถวที่ยังมี "เลขตู้" อยู่ในคอลัมน์ container เท่านั้น
    # (แถวข้อความท้ายตาราง เช่น "ขอแสดงความนับถือ," จะไม่ถูกแตะ)
    if tbl.get("clear_first"):
        col_indices = [column_index_from_string(c) for c in cols.values()]
        cmin, cmax = min(col_indices), max(col_indices)
        cn_col = column_index_from_string(cols["container"])
        r = start
        while r < start + max_rows:
            cn_val = ws.cell(row=r, column=cn_col).value
            if cn_val is None or (isinstance(cn_val, str) and not cn_val.strip()):
                break
            for c in range(cmin, cmax + 1):
                ws.cell(row=r, column=c).value = None
            r += 1

    # ---- เขียนข้อมูลตู้ทีละแถว ----
    written = 0
    for i, rec in enumerate(rows):
        excel_row = start + i
        if i >= max_rows:
            warnings.append(
                f"แม่แบบรองรับสูงสุด {max_rows} ตู้ ตู้ที่ {i + 1} เป็นต้นไปไม่ได้ถูกเขียน"
            )
            break

        for field, col_letter in cols.items():
            coord = _cell(col_letter, excel_row)
            if field in row_constants:              # ค่าคงที่ต่อแถว (เช่น terminal, payment)
                _set(ws, coord, row_constants[field])
            elif field == "no":
                _set(ws, coord, i + 1)
            elif field == "line":                   # คอลัมน์ LINE => HAL
                _set(ws, coord, LINE_CODE)
            elif field == "status":
                _set(ws, coord, status_code(rec["status"]))
            elif field in rec:
                _set(ws, coord, rec[field])

        written += 1

    # ---- บันทึกไฟล์ผลลัพธ์ ----
    os.makedirs(out_dir, exist_ok=True)
    out_path = os.path.join(out_dir, conf["out"])
    try:
        wb.save(out_path)
    except PermissionError:
        stamp = datetime.now().strftime("%Y%m%d-%H%M%S")
        stem, ext = os.path.splitext(conf["out"])
        out_path = os.path.join(out_dir, f"{stem}_{stamp}{ext}")
        wb.save(out_path)
        warnings.append("ไฟล์เดิมเปิดค้างอยู่ จึงบันทึกเป็นชื่อใหม่")

    return {
        "ok": True,
        "out_file": os.path.basename(out_path),
        "out_path": out_path,
        "rows": written,
        "terminal_key": conf["key"],
        "warnings": warnings,
        "error": "",
    }
