# -*- coding: utf-8 -*-
"""
xls_convert.py
แปลงไฟล์แม่แบบ .xls (Excel 97-2003) ให้เป็น .xlsx เพื่อให้ openpyxl แก้ไขได้
โดยพยายามรักษาหน้าตา/เลย์เอาต์เดิมไว้ให้มากที่สุด

ลำดับวิธีที่ใช้:
  1) Microsoft Excel ผ่าน COM (pywin32)  -> รักษาเลย์เอาต์ได้ครบที่สุด
  2) ถ้าไม่มี Excel/pywin32 -> อ่านค่าด้วย xlrd แล้วสร้าง .xlsx ใหม่ (ได้เฉพาะค่า ไม่มีการจัดรูปแบบ)

ผลลัพธ์ถูก cache ไว้ในโฟลเดอร์ build/templates/
"""

from __future__ import annotations

import os
import shutil

XLSX_FORMAT_CODE = 51  # xlOpenXMLWorkbook


def _excel_com_convert(src: str, dst: str) -> bool:
    try:
        import win32com.client as win32
    except Exception:
        return False

    excel = None
    try:
        excel = win32.DispatchEx("Excel.Application")
        excel.Visible = False
        excel.DisplayAlerts = False
        wb = excel.Workbooks.Open(os.path.abspath(src), ReadOnly=True)
        if os.path.exists(dst):
            os.remove(dst)
        wb.SaveAs(os.path.abspath(dst), FileFormat=XLSX_FORMAT_CODE)
        wb.Close(SaveChanges=False)
        return True
    except Exception as exc:  # pragma: no cover - ขึ้นกับเครื่อง
        print(f"  [xls_convert] Excel COM ล้มเหลว: {exc}")
        return False
    finally:
        if excel is not None:
            try:
                excel.Quit()
            except Exception:
                pass


def _xlrd_convert(src: str, dst: str) -> bool:
    try:
        import xlrd
        from openpyxl import Workbook
    except Exception as exc:
        print(f"  [xls_convert] ไม่มี xlrd/openpyxl: {exc}")
        return False

    try:
        book = xlrd.open_workbook(src)
    except Exception as exc:
        print(f"  [xls_convert] xlrd เปิดไฟล์ไม่ได้: {exc}")
        return False

    out = Workbook()
    out.remove(out.active)
    for sheet in book.sheets():
        ws = out.create_sheet(title=sheet.name[:31])
        for r in range(sheet.nrows):
            for c in range(sheet.ncols):
                val = sheet.cell_value(r, c)
                if val == "":
                    continue
                ctype = sheet.cell_type(r, c)
                if ctype == xlrd.XL_CELL_DATE:
                    try:
                        val = xlrd.xldate_as_datetime(val, book.datemode)
                    except Exception:
                        pass
                elif ctype == xlrd.XL_CELL_NUMBER and float(val).is_integer():
                    val = int(val)
                ws.cell(row=r + 1, column=c + 1, value=val)
    try:
        out.save(dst)
        return True
    except Exception as exc:
        print(f"  [xls_convert] บันทึก .xlsx ไม่ได้: {exc}")
        return False


def ensure_xlsx(src_path: str, cache_dir: str, force: bool = False) -> str:
    """
    รับ path ของไฟล์แม่แบบ (.xls หรือ .xlsx)
    คืน path ของไฟล์ .xlsx ที่พร้อมให้ openpyxl เปิด (อยู่ใน cache_dir)
    """
    os.makedirs(cache_dir, exist_ok=True)
    base = os.path.basename(src_path)
    stem, ext = os.path.splitext(base)
    ext = ext.lower()

    dst = os.path.join(cache_dir, stem + ".xlsx")

    if ext == ".xlsx":
        # เป็น .xlsx อยู่แล้ว - คัดลอกเข้ามาเป็นแม่แบบใน cache
        if force or not os.path.exists(dst) or os.path.getmtime(src_path) > os.path.getmtime(dst):
            shutil.copyfile(src_path, dst)
        return dst

    if not force and os.path.exists(dst) and os.path.getmtime(dst) >= os.path.getmtime(src_path):
        return dst

    print(f"  [xls_convert] กำลังแปลง {base} -> {os.path.basename(dst)}")
    if _excel_com_convert(src_path, dst):
        print("  [xls_convert] แปลงด้วย Microsoft Excel สำเร็จ (รักษาเลย์เอาต์)")
        return dst
    if _xlrd_convert(src_path, dst):
        print("  [xls_convert] แปลงด้วย xlrd สำเร็จ (เฉพาะค่า ไม่มีการจัดรูปแบบ)")
        return dst

    raise RuntimeError(
        f"แปลงไฟล์ {base} เป็น .xlsx ไม่สำเร็จ - "
        f"กรุณาเปิดไฟล์นี้ใน Excel แล้ว Save As เป็น .xlsx เองก่อน"
    )
