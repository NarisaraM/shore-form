# -*- coding: utf-8 -*-
"""
build_docs.py
สร้างไฟล์สำหรับเวอร์ชัน static (GitHub Pages) ไว้ในโฟลเดอร์ docs/

ทำสิ่งต่อไปนี้:
  1) แปลงไฟล์แม่แบบ .xls -> .xlsx (ผ่าน lib/xls_convert)
  2) คัดลอกเข้า docs/templates/ โดยตั้งชื่อไม่มีช่องว่าง
  3) "ลบข้อมูลตัวอย่าง/ข้อมูลส่วนบุคคล" (เบอร์โทรเจ้าหน้าที่, ตารางตัวอย่าง) ออกจากสำเนาใน docs/
     *ไฟล์แม่แบบต้นฉบับที่ราก repo ไม่ถูกแตะต้อง*
  4) คัดลอก Check.xlsx เข้า docs/
  5) วางไฟล์ .nojekyll

รันครั้งเดียว:  python build_docs.py
"""

from __future__ import annotations

import os
import shutil

import openpyxl

from lib.xls_convert import ensure_xlsx

BASE = os.path.dirname(os.path.abspath(__file__))
DOCS = os.path.join(BASE, "docs")
DOCS_TPL = os.path.join(DOCS, "templates")
CACHE = os.path.join(BASE, "build", "templates")

# ไฟล์แม่แบบต้นฉบับ -> ชื่อไฟล์ใน docs/templates/ (ไม่มีช่องว่าง)
TEMPLATES = {
    "A0-SHORE.xls": "A0-SHORE.xlsx",
    "B3-SHORE.xls": "B3-SHORE.xlsx",
    "B5C3-SHORE.xls": "B5C3-SHORE.xlsx",
    "A2-FORM  A.xlsx": "A2-FORM-A.xlsx",
    "A3C1C2-HUTCHISON.xls": "A3C1C2-HUTCHISON.xlsx",
}

# ช่วงเซลล์ที่ต้อง "ล้างค่า" ในสำเนา docs/  (sheet -> รายการ (col, row_from, row_to))
SCRUB = {
    "A0-SHORE.xlsx": {
        "Sheet1": [
            ("C", 29, 33),               # เบอร์โทรเจ้าหน้าที่ EAK/LEE/MINT/OIL/KUK
            *[(c, 30, 43) for c in "DEFGHIJK"],  # แถวตัวอย่างเที่ยวเรือเก่า
            ("G", 27, 27),               # โน้ตที่มีเบอร์มือถือเอเย่น
        ],
    },
    "B3-SHORE.xlsx": {
        "CHORE CY": [("A", 37, 37)],      # OOY 02-0596275
    },
    "B5C3-SHORE.xlsx": {
        "Sheet1": [(c, 2, 9) for c in "ABCDEFGHI"],   # ตารางตัวอย่าง
        "DataImport": [("A", 39, 43), ("A", 60, 70)], # เบอร์โทรเจ้าหน้าที่ + รายชื่อเรือตัวอย่าง
    },
    "A3C1C2-HUTCHISON.xlsx": {
        "Sheet1": [(c, 2, 9) for c in "ABCDEFGHI"],   # ตารางตัวอย่าง
        "HPT": [("A", 29, 29)],                       # โน้ตตัวอย่าง SHORE CHANGE VESSEL...
    },
    "A2-FORM-A.xlsx": {},
}


def clear_range(ws, col, r1, r2):
    n = 0
    for r in range(r1, r2 + 1):
        cell = ws[f"{col}{r}"]
        if cell.value not in (None, ""):
            cell.value = None
            n += 1
    return n


def strip_drawings(wb):
    """
    เอารูป/ภาพวาด/ชาร์ต ออกจากทุกชีต
    (ExcelJS อ่านไฟล์ที่มี drawing บางแบบไม่ได้ - error 'anchors')
    """
    for ws in wb.worksheets:
        try:
            ws._images = []
        except Exception:
            pass
        try:
            ws._charts = []
        except Exception:
            pass


def main():
    os.makedirs(DOCS_TPL, exist_ok=True)

    for src_name, out_name in TEMPLATES.items():
        src = os.path.join(BASE, src_name)
        if not os.path.isfile(src):
            print(f"  [!] ไม่พบ {src_name} - ข้าม")
            continue

        xlsx = ensure_xlsx(src, CACHE)
        dst = os.path.join(DOCS_TPL, out_name)

        # โหลดผ่าน openpyxl เสมอ เพื่อ (1) ตัด drawing/รูป ให้ ExcelJS อ่านได้
        # (2) ล้างข้อมูลตัวอย่าง/ส่วนบุคคล
        wb = openpyxl.load_workbook(xlsx)
        strip_drawings(wb)

        rules = SCRUB.get(out_name, {})
        total = 0
        for sheet, ranges in rules.items():
            if sheet not in wb.sheetnames:
                print(f"  [!] {out_name}: ไม่พบชีต {sheet!r}")
                continue
            ws = wb[sheet]
            for (col, r1, r2) in ranges:
                total += clear_range(ws, col, r1, r2)
        wb.save(dst)
        print(f"  {out_name}: ตัด drawing + ล้างข้อมูลตัวอย่าง/ส่วนบุคคล {total} เซลล์")

    # Check.xlsx
    shutil.copyfile(os.path.join(BASE, "Check.xlsx"), os.path.join(DOCS, "Check.xlsx"))
    print("  คัดลอก Check.xlsx")

    # .nojekyll (กัน GitHub Pages ประมวลผลแบบ Jekyll)
    open(os.path.join(DOCS, ".nojekyll"), "w").close()
    print("  สร้าง docs/.nojekyll")

    print("\nเสร็จแล้ว - เนื้อหา static อยู่ใน docs/")


if __name__ == "__main__":
    main()
